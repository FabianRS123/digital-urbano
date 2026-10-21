"""Exportación de figuras (PNG, SVG, PDF) y de tablas (CSV, Excel, texto para copiar y pegar)."""

from __future__ import annotations

import io
import zipfile
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .catalogo import FIGURAS, Figura


def tablas_de(datos: Any) -> dict[str, pd.DataFrame]:
    if isinstance(datos, pd.DataFrame):
        return {"Datos": datos}
    return dict(datos)


# ------------------------------------------------------------------ figuras

def png(figura: Figura, fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=figura.dpi, facecolor="white")
    return buf.getvalue()


def svg(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="svg", facecolor="white")
    return buf.getvalue()


def pdf(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="pdf", facecolor="white")
    return buf.getvalue()


# ------------------------------------------------------------------ tablas

def csv_bytes(tabla: pd.DataFrame) -> bytes:
    """CSV UTF-8 con BOM para que Excel respete los acentos."""
    return tabla.to_csv(index=False).encode("utf-8-sig")


def _numero(valor: Any) -> str:
    if isinstance(valor, float):
        if valor != valor:
            return ""
        texto = f"{valor:.6g}" if abs(valor) < 1e6 else f"{valor:.0f}"
        return texto.replace(".", ",")
    return "" if valor is None else str(valor)


def texto_para_copiar(tabla: pd.DataFrame) -> str:
    """Texto separado por tabuladores y con coma decimal: se pega como tabla en Excel, Word o Google Sheets."""
    encabezado = "\t".join(str(c) for c in tabla.columns)
    filas = ("\t".join(_numero(v) for v in fila) for fila in tabla.itertuples(index=False, name=None))
    return "\n".join([encabezado, *filas])


def excel_bytes(tablas: dict[str, pd.DataFrame]) -> bytes:
    """Libro de Excel con una hoja por tabla, encabezado resaltado y columnas ajustadas."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as escritor:
        for nombre, tabla in tablas.items():
            hoja = "".join("-" if c in "[]:*?/\\" else c for c in nombre)[:31]
            tabla.to_excel(escritor, sheet_name=hoja, index=False)
            ws = escritor.sheets[hoja]
            for celda in ws[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="1D4F91")
                celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, columna in enumerate(tabla.columns, start=1):
                ancho = max([len(str(columna))] + [len(str(v)) for v in tabla.iloc[:200, i - 1]]) + 3
                ws.column_dimensions[get_column_letter(i)].width = min(max(ancho, 10), 60)
            ws.freeze_panes = "A2"
    return buf.getvalue()


# ------------------------------------------------------------------ todo en un ZIP

def todo(construidas: dict[str, tuple[Any, Any]], comprobaciones: pd.DataFrame | None = None, sufijo: str = "") -> bytes:
    """ZIP con PNG, SVG, PDF y datos de todas las figuras, más la tabla de comprobaciones."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in FIGURAS:
            fig, datos = construidas[f.clave]
            base = f.nombre_archivo + sufijo
            zf.writestr(f"{base}.png", png(f, fig))
            zf.writestr(f"{base}.svg", svg(fig))
            zf.writestr(f"{base}.pdf", pdf(fig))
            zf.writestr(f"datos/{base}.xlsx", excel_bytes(tablas_de(datos)))
        if comprobaciones is not None:
            zf.writestr("comprobaciones.csv", csv_bytes(comprobaciones))
    return buf.getvalue()
